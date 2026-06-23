from __future__ import annotations

import json
import os
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from teacher_agent.schemas import ClassProfile, OutputBundle, ensure_output_dir, slugify


ASSET_MANIFEST_PATH = Path("assets/ppt_images/manifest.json")
BEAMER_FORMAT_PATH = Path("templates/beamer/format.tex")
LATEX_COMPILE_TIMEOUT_SECONDS = 60


@dataclass
class PPTRequest:
    topic: str
    class_profile: ClassProfile
    learning_objectives: list[str] = field(default_factory=list)
    slide_count: int = 10
    output_dir: str | Path = "outputs"
    compile_pdf: bool = True


@dataclass
class SlideSpec:
    title: str
    bullets: list[str]
    speaker_notes: str = ""
    visual_key: str = ""


def build_slide_plan(request: PPTRequest) -> list[SlideSpec]:
    weak_points = request.class_profile.weak_points
    objectives = request.learning_objectives or [
        f"理解 {request.topic} 的核心概念",
        "能够解释方法适用条件",
        "能够根据真实问题选择合适的统计方法",
    ]

    slides = [
        SlideSpec(
            title=f"{request.topic}",
            bullets=[
                request.class_profile.course_name,
                f"面向：{request.class_profile.class_name}",
                "基于班级学习画像动态调整",
            ],
            speaker_notes="开场时说明本节课会重点照顾班级薄弱点。",
            visual_key=_default_visual_for_topic(request.topic),
        ),
        SlideSpec(
            title="本节学习目标",
            bullets=objectives,
        ),
        SlideSpec(
            title="班级学习画像反馈",
            bullets=_profile_bullets(request.class_profile),
            speaker_notes="这里用于把后续讲解和学生真实困难连接起来。",
        ),
    ]

    slides.extend(_topic_specific_slides(request.topic))

    for item in weak_points:
        bullets = [
            f"当前掌握度：{_format_mastery(item.mastery)}",
            f"优先级：{item.priority}",
        ]
        bullets.extend(f"常见误区：{error}" for error in item.common_errors[:3])
        bullets.append("教学处理：先用直觉例子解释，再回到公式和适用条件。")
        slides.append(
            SlideSpec(
                title=f"重点补强：{item.name}",
                bullets=bullets,
                speaker_notes=f"围绕 {item.name} 安排一次即时提问或小练习。",
            )
        )

    slides.extend(
        [
            SlideSpec(
                title="核心概念框架",
                bullets=[
                    "研究问题：变量、总体、样本和假设",
                    "统计模型：假设、参数、估计和不确定性",
                    "解释结果：效应量、置信区间和实际意义",
                ],
            ),
            SlideSpec(
                title="课堂例题",
                bullets=[
                    f"围绕 {request.topic} 设计一个真实生物医学数据情境",
                    "先让学生判断变量类型和研究问题",
                    "再选择统计方法并解释结果",
                ],
            ),
            SlideSpec(
                title="即时练习",
                bullets=[
                    "1 道概念判断题",
                    "1 道方法选择题",
                    "1 道结果解释题",
                ],
            ),
            SlideSpec(
                title="课后巩固",
                bullets=[
                    "复盘今天最容易混淆的概念",
                    "完成针对薄弱点的短练习",
                    "把错题原因反馈给学习画像",
                ],
            ),
        ]
    )

    return slides[: max(4, request.slide_count)]


def _topic_specific_slides(topic: str) -> list[SlideSpec]:
    normalized = topic.lower()
    if "logistic" in normalized or "逻辑" in topic or "logit" in normalized:
        return [
            SlideSpec(
                title="Logistic 回归要解决的问题",
                bullets=[
                    "结局变量是二分类或事件发生/不发生",
                    "模型输出的是对数优势，而不是直接概率差",
                    "适合分析风险因素、诊断指标和疾病结局",
                ],
                visual_key="logistic_curve",
            ),
            SlideSpec(
                title="比值比的解释",
                bullets=[
                    "OR > 1 表示暴露组优势增加",
                    "OR < 1 表示暴露组优势降低",
                    "解释时要同时看置信区间和研究设计",
                ],
            ),
            SlideSpec(
                title="学生常见误区",
                bullets=[
                    "把 OR 直接解释成概率差",
                    "只看 p 值，不看效应大小",
                    "忽略混杂因素和模型诊断",
                ],
            ),
        ]
    if "多重" in topic or "fdr" in normalized or "fwer" in normalized:
        return [
            SlideSpec(
                title="为什么会有多重检验问题",
                bullets=[
                    "检验次数越多，偶然显著的概率越高",
                    "基因组数据常同时检验成千上万个基因",
                    "不校正会夸大发现数量",
                ],
            ),
            SlideSpec(
                title="FWER 与 FDR",
                bullets=[
                    "FWER 控制至少一个假阳性的概率",
                    "FDR 控制发现集合中的假阳性比例",
                    "高通量筛选中常优先考虑 FDR",
                ],
            ),
            SlideSpec(
                title="Benjamini-Hochberg 思路",
                bullets=[
                    "将 p 值从小到大排序",
                    "与逐步放宽的阈值比较",
                    "找到可接受的发现集合",
                ],
                visual_key="benjamini_hochberg",
            ),
        ]
    if "配对" in topic or "paired" in normalized:
        return [
            SlideSpec(
                title="配对设计的核心",
                bullets=[
                    "同一对象或匹配对象形成一对观测",
                    "分析对象是每一对的差值",
                    "配对能减少个体差异带来的噪声",
                ],
            ),
            SlideSpec(
                title="配对 t 检验条件",
                bullets=[
                    "差值近似来自正态分布",
                    "配对关系来自设计而不是事后强行匹配",
                    "异常差值需要结合背景判断",
                ],
            ),
            SlideSpec(
                title="结果报告",
                bullets=[
                    "报告平均差值和置信区间",
                    "解释方向和实际意义",
                    "说明样本量和检验假设",
                ],
            ),
        ]
    return [
        SlideSpec(
            title=f"{topic} 的问题框架",
            bullets=[
                "先明确研究问题和变量类型",
                "再判断统计方法的适用条件",
                "最后解释结果的不确定性和实际意义",
            ],
        )
    ]


def generate_ppt(request: PPTRequest) -> OutputBundle:
    output_dir = ensure_output_dir(request.output_dir)
    base_name = slugify(f"{request.topic}_课件")
    slides = build_slide_plan(request)

    tex_path = output_dir / f"{base_name}.tex"
    tex_path.write_text(_render_beamer_latex(request, slides), encoding="utf-8")

    should_compile = request.compile_pdf or os.getenv("AGENT_COMPILE_LATEX", "").lower() in {
        "1",
        "true",
        "yes",
        "on",
    }
    pdf_path = _try_compile_beamer(tex_path) if should_compile else None
    pptx_path = _render_designed_pptx(output_dir / f"{base_name}.pptx", request, slides)
    return OutputBundle(
        source_path=tex_path,
        artifact_path=pptx_path,
        metadata={
            "topic": request.topic,
            "slide_count": len(slides),
            "source_format": "latex",
            "pdf_compile_enabled": should_compile,
            "pdf_available": pdf_path is not None,
            "pptx_available": pptx_path is not None,
        },
    )


def _profile_bullets(profile: ClassProfile) -> list[str]:
    bullets = []
    if profile.weak_points:
        weak = "、".join(item.name for item in profile.weak_points[:4])
        bullets.append(f"主要薄弱点：{weak}")
    if profile.strong_points:
        bullets.append(f"优势基础：{'、'.join(profile.strong_points[:4])}")
    bullets.extend(profile.teacher_notes[:3])
    return bullets or ["暂无画像数据，先按章节基础目标授课。"]


def _format_mastery(value: float | None) -> str:
    if value is None:
        return "未知"
    return f"{value:.0%}"


def _render_markdown(request: PPTRequest, slides: list[SlideSpec]) -> str:
    lines = [
        f"# {request.topic} 课件",
        "",
        f"- 课程：{request.class_profile.course_name}",
        f"- 班级：{request.class_profile.class_name}",
        f"- 幻灯片数：{len(slides)}",
        "",
    ]
    for index, slide in enumerate(slides, start=1):
        lines.extend([f"## {index}. {slide.title}", ""])
        lines.extend(f"- {bullet}" for bullet in slide.bullets)
        if slide.speaker_notes:
            lines.extend(["", f"> 讲者备注：{slide.speaker_notes}"])
        lines.append("")
    return "\n".join(lines)


def _render_beamer_latex(request: PPTRequest, slides: list[SlideSpec]) -> str:
    manifest = _load_asset_manifest()
    title = _latex_escape(request.topic)
    course = _latex_escape(request.class_profile.course_name)
    class_name = _latex_escape(request.class_profile.class_name)
    format_path = f"../{BEAMER_FORMAT_PATH.as_posix()}"

    lines = [
        r"\documentclass[aspectratio=169,UTF8,fontset=none]{ctexbeamer}",
        rf"\input{{{format_path}}}",
        rf"\title[{title}]{{{title}}}",
        rf"\subtitle{{{course} · {class_name}}}",
        r"\author{AI 教学 Agent}",
        r"\date{\today}",
        "",
        r"\begin{document}",
        "",
        r"\begin{frame}[plain]",
        r"  \titlepage",
        r"\end{frame}",
        "",
    ]

    for index, slide in enumerate(slides[1:], start=2):
        asset = _asset_for_slide(slide, manifest)
        lines.extend(_render_slide_frame(index, slide, asset))

    lines.extend([r"\end{document}", ""])
    return "\n".join(lines)


def _render_slide_frame(
    index: int,
    slide: SlideSpec,
    asset: dict[str, Any] | None,
) -> list[str]:
    title = _latex_escape(slide.title)
    kicker = "BIOSTATISTICS TEACHING AGENT"
    if asset:
        return [
            rf"\begin{{frame}}[t]{{{title}}}",
            rf"  \AgentKicker{{{kicker}}}",
            r"  \begin{columns}[T,onlytextwidth]",
            r"    \begin{column}{0.56\textwidth}",
            r"      \begin{itemize}",
            *[rf"        \item {_latex_escape(bullet)}" for bullet in slide.bullets],
            r"      \end{itemize}",
            r"    \end{column}",
            r"    \begin{column}{0.40\textwidth}",
            rf"      \includegraphics[width=\linewidth,height=0.56\textheight,keepaspectratio]{{{asset.get('latex_path', asset['path'])}}}",
            rf"      \AgentImageCredit{{{_latex_escape(asset.get('title', '课程配图'))} · {_latex_escape(asset.get('source', 'Wikimedia Commons'))}}}",
            r"    \end{column}",
            r"  \end{columns}",
            r"\end{frame}",
            "",
        ]

    return [
        rf"\begin{{frame}}[t]{{{title}}}",
        rf"  \AgentKicker{{{kicker}}}",
        r"  \begin{itemize}",
        *[rf"    \item {_latex_escape(bullet)}" for bullet in slide.bullets],
        r"  \end{itemize}",
        _speaker_note_callout(slide),
        r"\end{frame}",
        "",
    ]


def _speaker_note_callout(slide: SlideSpec) -> str:
    if not slide.speaker_notes:
        return ""
    return rf"  \AgentCallout{{讲者提示}}{{{_latex_escape(slide.speaker_notes)}}}"


def _load_asset_manifest() -> dict[str, Any]:
    if not ASSET_MANIFEST_PATH.exists():
        return {}
    try:
        data = json.loads(ASSET_MANIFEST_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}
    return data.get("assets", {})


def _asset_for_slide(slide: SlideSpec, manifest: dict[str, Any]) -> dict[str, Any] | None:
    if not slide.visual_key:
        return None
    asset = manifest.get(slide.visual_key)
    if not asset:
        return None
    path = Path(asset.get("path", ""))
    if not path.exists():
        return None
    return {
        "path": path.as_posix(),
        "latex_path": f"../{path.as_posix()}",
        "title": asset.get("title") or slide.visual_key,
        "source": "Wikimedia Commons",
    }


def _default_visual_for_topic(topic: str) -> str:
    normalized = topic.lower()
    if "logistic" in normalized or "逻辑" in topic or "logit" in normalized:
        return "logistic_curve"
    if "多重" in topic or "fdr" in normalized or "fwer" in normalized:
        return "benjamini_hochberg"
    if "基因" in topic or "genome" in normalized or "dna" in normalized:
        return "dna_double_helix"
    return "dna_double_helix"


def _try_compile_beamer(tex_path: Path) -> Path | None:
    tectonic = shutil.which("tectonic")
    compiler = shutil.which("xelatex") or shutil.which("pdflatex")
    if tectonic:
        command = [tectonic, "-o", str(tex_path.parent), str(tex_path)]
    elif compiler:
        command = [
            compiler,
            "-interaction=nonstopmode",
            "-halt-on-error",
            "-output-directory",
            str(tex_path.parent),
            str(tex_path),
        ]
    else:
        return None

    try:
        subprocess.run(
            command,
            check=True,
            capture_output=True,
            text=True,
            timeout=LATEX_COMPILE_TIMEOUT_SECONDS,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return None
    return tex_path.with_suffix(".pdf")


def _render_designed_pptx(
    path: Path,
    request: PPTRequest,
    slides: list[SlideSpec],
) -> Path | None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        return None

    manifest = _load_asset_manifest()
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    colors = {
        "ink": RGBColor(23, 32, 42),
        "muted": RGBColor(102, 112, 133),
        "blue": RGBColor(31, 111, 235),
        "teal": RGBColor(14, 124, 134),
        "sand": RGBColor(245, 247, 250),
        "white": RGBColor(255, 255, 255),
    }

    for index, slide in enumerate(slides):
        ppt_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _paint_background(ppt_slide, presentation, colors)
        if index == 0:
            _render_pptx_title_slide(ppt_slide, presentation, request, slide, colors)
            continue
        asset = _asset_for_slide(slide, manifest)
        _render_pptx_content_slide(
            ppt_slide,
            presentation,
            slide,
            colors,
            index=index,
            total=len(slides),
            asset=asset,
        )

    presentation.save(path)
    return path


def _paint_background(slide, presentation, colors: dict[str, Any]) -> None:
    from pptx.util import Inches

    width = presentation.slide_width
    height = presentation.slide_height
    bg = slide.shapes.add_shape(1, 0, 0, width, height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = colors["white"]
    bg.line.fill.background()

    band = slide.shapes.add_shape(1, 0, 0, width, Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = colors["blue"]
    band.line.fill.background()


def _render_pptx_title_slide(slide, presentation, request: PPTRequest, spec: SlideSpec, colors) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    accent = slide.shapes.add_shape(1, Inches(0.65), Inches(0.85), Inches(0.15), Inches(5.8))
    accent.fill.solid()
    accent.fill.fore_color.rgb = colors["teal"]
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(1.05), Inches(1.2), Inches(7.3), Inches(1.2))
    title = title_box.text_frame.paragraphs[0]
    title.text = request.topic
    title.font.name = "Microsoft YaHei"
    title.font.size = Pt(42)
    title.font.bold = True
    title.font.color.rgb = colors["ink"]

    subtitle_box = slide.shapes.add_textbox(Inches(1.08), Inches(2.55), Inches(7.8), Inches(0.6))
    subtitle = subtitle_box.text_frame.paragraphs[0]
    subtitle.text = f"{request.class_profile.course_name} · {request.class_profile.class_name}"
    subtitle.font.name = "Microsoft YaHei"
    subtitle.font.size = Pt(18)
    subtitle.font.color.rgb = colors["muted"]

    bullets_box = slide.shapes.add_textbox(Inches(1.08), Inches(3.35), Inches(6.3), Inches(1.5))
    frame = bullets_box.text_frame
    frame.clear()
    for idx, bullet_text in enumerate(spec.bullets):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = bullet_text
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = colors["ink"]

    asset = _asset_for_slide(spec, _load_asset_manifest())
    if asset:
        slide.shapes.add_picture(
            asset["path"],
            Inches(8.25),
            Inches(1.25),
            width=Inches(4.3),
            height=Inches(4.0),
        )

    footer = slide.shapes.add_textbox(Inches(1.08), Inches(6.55), Inches(11.2), Inches(0.3))
    p = footer.text_frame.paragraphs[0]
    p.text = "AI 教学 Agent · LaTeX/PPTX 双格式课件"
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(11)
    p.font.color.rgb = colors["muted"]


def _render_pptx_content_slide(
    slide,
    presentation,
    spec: SlideSpec,
    colors,
    *,
    index: int,
    total: int,
    asset: dict[str, Any] | None,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(8.8), Inches(0.55))
    title = title_box.text_frame.paragraphs[0]
    title.text = spec.title
    title.font.name = "Microsoft YaHei"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = colors["ink"]

    kicker_box = slide.shapes.add_textbox(Inches(0.78), Inches(1.12), Inches(4.5), Inches(0.25))
    kicker = kicker_box.text_frame.paragraphs[0]
    kicker.text = "BIOSTATISTICS TEACHING AGENT"
    kicker.font.name = "Arial"
    kicker.font.size = Pt(8)
    kicker.font.bold = True
    kicker.font.color.rgb = colors["teal"]

    bullet_width = Inches(6.6 if asset else 11.6)
    bullets_box = slide.shapes.add_textbox(Inches(0.85), Inches(1.65), bullet_width, Inches(4.7))
    frame = bullets_box.text_frame
    frame.clear()
    frame.word_wrap = True
    for bullet_index, bullet_text in enumerate(spec.bullets):
        paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
        paragraph.text = bullet_text
        paragraph.level = 0
        paragraph.space_after = Pt(10)
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(19)
        paragraph.font.color.rgb = colors["ink"]

    if asset:
        panel = slide.shapes.add_shape(1, Inches(8.0), Inches(1.55), Inches(4.55), Inches(4.4))
        panel.fill.solid()
        panel.fill.fore_color.rgb = colors["sand"]
        panel.line.color.rgb = RGBColor(216, 222, 233)
        slide.shapes.add_picture(
            asset["path"],
            Inches(8.25),
            Inches(1.82),
            width=Inches(4.05),
            height=Inches(3.25),
        )
        credit = slide.shapes.add_textbox(Inches(8.25), Inches(5.25), Inches(4.05), Inches(0.45))
        credit_p = credit.text_frame.paragraphs[0]
        credit_p.text = f"{asset.get('title', '课程配图')} · Wikimedia Commons"
        credit_p.font.name = "Arial"
        credit_p.font.size = Pt(8)
        credit_p.font.color.rgb = colors["muted"]

    if spec.speaker_notes:
        note = slide.shapes.add_textbox(Inches(0.85), Inches(6.2), Inches(9.0), Inches(0.45))
        p = note.text_frame.paragraphs[0]
        p.text = f"讲者提示：{spec.speaker_notes}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(11)
        p.font.color.rgb = colors["muted"]

    page = slide.shapes.add_textbox(Inches(11.55), Inches(6.8), Inches(1.0), Inches(0.25))
    p = page.text_frame.paragraphs[0]
    p.text = f"{index + 1}/{total}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = colors["muted"]


def _latex_escape(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(char, char) for char in text)


def _try_render_pptx(path: Path, slides: list[SlideSpec]) -> Path | None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        return None

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for index, slide in enumerate(slides):
        layout = presentation.slide_layouts[0] if index == 0 else presentation.slide_layouts[1]
        ppt_slide = presentation.slides.add_slide(layout)
        ppt_slide.shapes.title.text = slide.title

        placeholder = ppt_slide.placeholders[1]
        text_frame = placeholder.text_frame
        text_frame.clear()
        for bullet_index, bullet in enumerate(slide.bullets):
            paragraph = text_frame.paragraphs[0] if bullet_index == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(24 if index == 0 else 20)
            paragraph.font.name = "Microsoft YaHei"

        if slide.speaker_notes:
            notes = ppt_slide.notes_slide.notes_text_frame
            notes.text = slide.speaker_notes

    presentation.save(path)
    return path
